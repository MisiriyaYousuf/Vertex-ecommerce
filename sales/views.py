from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Sum, Count, F, DecimalField, ExpressionWrapper
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.utils import timezone
from django.views.decorators.cache import never_cache
from orders.models import Order, OrderItem
from .forms import SalesReportForm
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def get_report_date_range(form):

    cleaned_data = form.cleaned_data

    report_type = cleaned_data["report_type"]

    today = timezone.localdate()

    if report_type == "daily":

        selected_date = cleaned_data["date"]

        return (
            selected_date,
            selected_date,
            "Daily Report"
        )

    elif report_type == "weekly":

        selected_date = cleaned_data["date"]

        start_date = (
            selected_date
            - timedelta(
                days=selected_date.weekday()
            )
        )

        end_date = start_date + timedelta(days=6)

        return (
            start_date,
            end_date,
            "Weekly Report"
        )

    elif report_type == "monthly":

        month_value = cleaned_data["month"]

        year, month = map(
            int,
            month_value.split("-")
        )

        start_date = datetime(
            year,
            month,
            1
        ).date()

        if month == 12:

            end_date = datetime(
                year + 1,
                1,
                1
            ).date() - timedelta(days=1)

        else:

            end_date = datetime(
                year,
                month + 1,
                1
            ).date() - timedelta(days=1)

        return (
            start_date,
            end_date,
            "Monthly Report"
        )

    elif report_type == "yearly":

        year = cleaned_data["year"]

        start_date = datetime(
            year,
            1,
            1
        ).date()

        end_date = datetime(
            year,
            12,
            31
        ).date()

        return (
            start_date,
            end_date,
            "Yearly Report"
        )

    elif report_type == "custom":

        return (
            cleaned_data["start_date"],
            cleaned_data["end_date"],
            "Custom Date Report"
        )

    return today, today, "Sales Report"

def get_sales_orders(start_date, end_date):

    end_datetime = datetime.combine(
        end_date,
        datetime.max.time()
    )

    start_datetime = datetime.combine(
        start_date,
        datetime.min.time()
    )

    start_datetime = timezone.make_aware(
        start_datetime
    )

    end_datetime = timezone.make_aware(
        end_datetime
    )

    return (
        Order.objects
        .filter(
            created_at__gte=start_datetime,
            created_at__lte=end_datetime,
        )
        .exclude(
            status="Cancelled"
        )
        .select_related(
            "user"
        )
        .prefetch_related(
            "items"
        )
        .order_by(
            "-created_at"
        )
    )

@never_cache
@login_required
def sales_report(request):

    if not request.user.is_superuser:
        return redirect("users:home")

    form = SalesReportForm(
        request.GET or None
    )

    orders = Order.objects.none()

    context = {
        "form": form,
        "orders": orders,
        "report_generated": False,

        "start_date": None,
        "end_date": None,

        "sales_count": 0,
        "order_amount": Decimal("0.00"),

        "product_discount": Decimal("0.00"),
        "coupon_discount": Decimal("0.00"),
        "overall_discount": Decimal("0.00"),

        "net_sales": Decimal("0.00"),

        "report_title": "Sales Report",
    }

    if request.GET:

        if form.is_valid():

            (
                start_date,
                end_date,
                report_title
            ) = get_report_date_range(form)

            orders = get_sales_orders(
                start_date,
                end_date
            )

            sales_count = orders.count()

            order_amount = (
                orders.aggregate(
                    total=Sum("total_amount")
                )["total"]
                or Decimal("0.00")
            )

            product_discount = (
                orders.aggregate(
                    total=Sum(
                        "items__discount"
                    )
                )["total"]
                or Decimal("0.00")
            )

            coupon_discount = Decimal("0.00")

            for order in orders:

                coupon_discount += (
                    getattr(
                        order,
                        "coupon_discount",
                        Decimal("0.00")
                    )
                    or Decimal("0.00")
                )

            overall_discount = (
                product_discount
                + coupon_discount
            )

            net_sales = order_amount

            context.update({

                "orders": orders,

                "report_generated": True,

                "start_date": start_date,
                "end_date": end_date,

                "sales_count": sales_count,

                "order_amount": order_amount,

                "product_discount": product_discount,

                "coupon_discount": coupon_discount,

                "overall_discount": overall_discount,

                "net_sales": net_sales,

                "report_title": report_title,

            })

    return render(
        request,
        "sales_report.html",
        context
    )

@login_required
def sales_report_pdf(request):

    if not request.user.is_superuser:
        return redirect("users:home")

    form = SalesReportForm(
        request.GET or None
    )

    if not form.is_valid():
        messages.error(
            request,
            "Invalid report filters."
        )
        return redirect(
            "sales:sales_report"
        )

    (
        start_date,
        end_date,
        report_title
    ) = get_report_date_range(form)

    orders = get_sales_orders(
        start_date,
        end_date
    )

    sales_count = orders.count()

    order_amount = (
        orders.aggregate(
            total=Sum("total_amount")
        )["total"]
        or Decimal("0.00")
    )

    product_discount = (
        orders.aggregate(
            total=Sum("items__discount")
        )["total"]
        or Decimal("0.00")
    )

    coupon_discount = Decimal("0.00")

    for order in orders:

        coupon_discount += (
            getattr(
                order,
                "coupon_discount",
                Decimal("0.00")
            )
            or Decimal("0.00")
        )

    overall_discount = (
        product_discount
        + coupon_discount
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="sales_report.pdf"'
    )

    document = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "Vertex Sales Report",
            styles["Title"]
        )
    )

    elements.append(
        Spacer(1, 10)
    )

    elements.append(
        Paragraph(
            f"{report_title}",
            styles["Heading2"]
        )
    )

    elements.append(
        Paragraph(
            f"From: {start_date} &nbsp;&nbsp; "
            f"To: {end_date}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(1, 20)
    )

    summary_data = [
        ["Metric", "Amount / Count"],

        [
            "Overall Sales Count",
            str(sales_count)
        ],

        [
            "Overall Order Amount",
            f"₹ {order_amount:,.2f}"
        ],

        [
            "Product Discount",
            f"₹ {product_discount:,.2f}"
        ],

        [
            "Coupon Discount",
            f"₹ {coupon_discount:,.2f}"
        ],

        [
            "Overall Discount",
            f"₹ {overall_discount:,.2f}"
        ],

        [
            "Net Sales",
            f"₹ {order_amount:,.2f}"
        ],
    ]

    table = Table(
        summary_data,
        colWidths=[250, 200]
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#2c3e50")
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    elements.append(table)

    elements.append(
        Spacer(1, 25)
    )

    order_data = [
        [
            "Order",
            "Customer",
            "Date",
            "Payment",
            "Amount",
        ]
    ]

    for order in orders:

        customer = (
            order.user.username
            if order.user
            else "-"
        )

        order_data.append([
            f"#{order.id}",
            customer,
            order.created_at.strftime(
                "%d-%m-%Y"
            ),
            getattr(
                order,
                "payment_method",
                "-"
            ),
            f"₹ {order.total_amount:,.2f}",
        ])

    order_table = Table(
        order_data,
        repeatRows=1
    )

    order_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#394f66")
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    elements.append(order_table)

    document.build(elements)

    return response

@login_required
def sales_report_excel(request):

    if not request.user.is_superuser:
        return redirect("users:home")

    form = SalesReportForm(
        request.GET or None
    )

    if not form.is_valid():
        messages.error(
            request,
            "Invalid report filters."
        )
        return redirect(
            "sales:sales_report"
        )

    (
        start_date,
        end_date,
        report_title
    ) = get_report_date_range(form)

    orders = get_sales_orders(
        start_date,
        end_date
    )

    sales_count = orders.count()

    order_amount = (
        orders.aggregate(
            total=Sum("total_amount")
        )["total"]
        or Decimal("0.00")
    )

    product_discount = (
        orders.aggregate(
            total=Sum("items__discount")
        )["total"]
        or Decimal("0.00")
    )

    coupon_discount = Decimal("0.00")

    for order in orders:

        coupon_discount += (
            getattr(
                order,
                "coupon_discount",
                Decimal("0.00")
            )
            or Decimal("0.00")
        )

    overall_discount = (
        product_discount
        + coupon_discount
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Sales Report"

    worksheet["A1"] = "Vertex Sales Report"

    worksheet["A1"].font = Font(
        bold=True,
        size=18
    )

    worksheet["A2"] = report_title

    worksheet["A3"] = (
        f"From: {start_date} "
        f"To: {end_date}"
    )

    
    worksheet["A5"] = "Summary"

    worksheet["A5"].font = Font(
        bold=True,
        size=14
    )

    summary = [
        (
            "Overall Sales Count",
            sales_count
        ),

        (
            "Overall Order Amount",
            float(order_amount)
        ),

        (
            "Product Discount",
            float(product_discount)
        ),

        (
            "Coupon Discount",
            float(coupon_discount)
        ),

        (
            "Overall Discount",
            float(overall_discount)
        ),

        (
            "Net Sales",
            float(order_amount)
        ),
    ]

    row = 6

    for label, value in summary:

        worksheet.cell(
            row=row,
            column=1,
            value=label
        )

        worksheet.cell(
            row=row,
            column=2,
            value=value
        )

        row += 1

    row += 2

    headers = [
        "Order ID",
        "Customer",
        "Date",
        "Payment Method",
        "Payment Status",
        "Order Status",
        "Amount",
    ]

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=row,
            column=column,
            value=header
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2C3E50"
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    row += 1

    for order in orders:

        worksheet.cell(
            row=row,
            column=1,
            value=order.id
        )

        worksheet.cell(
            row=row,
            column=2,
            value=(
                order.user.username
                if order.user
                else "-"
            )
        )

        worksheet.cell(
            row=row,
            column=3,
            value=order.created_at.strftime(
                "%d-%m-%Y"
            )
        )

        worksheet.cell(
            row=row,
            column=4,
            value=getattr(
                order,
                "payment_method",
                "-"
            )
        )

        worksheet.cell(
            row=row,
            column=5,
            value=getattr(
                order,
                "payment_status",
                "-"
            )
        )

        worksheet.cell(
            row=row,
            column=6,
            value=getattr(
                order,
                "status",
                "-"
            )
        )

        worksheet.cell(
            row=row,
            column=7,
            value=float(
                order.total_amount
            )
        )

        row += 1

    widths = {
        "A": 18,
        "B": 25,
        "C": 15,
        "D": 20,
        "E": 20,
        "F": 18,
        "G": 18,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="sales_report.xlsx"'
    )

    return response